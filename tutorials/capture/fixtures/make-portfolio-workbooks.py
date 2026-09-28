#!/usr/bin/env python3
"""Build the tutorial-08 portfolio workbooks from the app's own template.

    python3 tutorials/capture/fixtures/make-portfolio-workbooks.py <portfolio-import-template.xlsx>

Download the template from the running app first (Properties & Units > More >
Import Portfolio > Download Template). Writes, next to this script:
  oasis-crest-garden-villas.xlsx          valid: 1 project, 2 buildings, 4 units, 2 tenants
  oasis-crest-garden-villas-invalid.xlsx  the same with two deliberate mistakes
All data is fictional (Oasis Crest demo brand, reserved .example addresses).
"""
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
PROJECT = "Oasis Crest Garden Villas"
PROPERTIES = [[PROJECT, "فلل حدائق القمة", "DUBAI", "Garden Crescent, Al Furjan, Dubai", "RESIDENTIAL", "40218-65173"]]
UNITS = [
    [PROJECT, "Villa Court 1", "GV-101", "BHK2", 1180, 92000],
    [PROJECT, "Villa Court 1", "GV-102", "BHK3", 1650, 125000],
    [PROJECT, "Villa Court 2", "GV-201", "BHK1", 760, 68000],
    [PROJECT, "Villa Court 2", "GV-202", "STUDIO", 480, 45000],
]
TENANTS = [
    ["Omar Al Mansoori", "عمر المنصوري", "omar.almansoori@oasiscrest.example", "+971500000731"],
    ["Leila Farouk", "ليلى فاروق", "leila.farouk@oasiscrest.example", "+971500000732"],
]


def fill(ws, rows):
    ws.delete_rows(2, ws.max_row)
    for row in rows:
        ws.append(row)


def build(template, out, units):
    wb = openpyxl.load_workbook(template)
    fill(wb["Properties"], PROPERTIES)
    fill(wb["Units"], units)
    fill(wb["Renters"], TENANTS)
    fill(wb["Leases"], [])
    if "Cheques" in wb.sheetnames:
        fill(wb["Cheques"], [])
    wb.save(out)
    print(out)


def main():
    template = sys.argv[1]
    build(template, HERE / "oasis-crest-garden-villas.xlsx", UNITS)
    invalid = [list(u) for u in UNITS]
    invalid[1][3] = "3 BED"                    # not one of the template's unit types
    invalid[3][0] = "Oasis Crest Garden Vilas"  # a project name the Properties sheet does not have
    build(template, HERE / "oasis-crest-garden-villas-invalid.xlsx", invalid)


if __name__ == "__main__":
    main()
